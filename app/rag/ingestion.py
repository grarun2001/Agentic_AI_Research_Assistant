import os
import json
import uuid
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional

import fitz  # PyMuPDF
import docx
import openpyxl
import pptx
import csv
from PIL import Image
try:
    import pytesseract
except ImportError:
    pytesseract = None

from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings
from app.rag.embeddings import embed_documents
from app.rag.milvus_store import milvus_store

logger = logging.getLogger(__name__)

# Supported document extensions
SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".csv", ".xlsx", ".pptx",
    ".png", ".jpg", ".jpeg", ".webp"
}

def load_chunks_index() -> List[Dict[str, Any]]:
    """Loads all persisted chunks from data/index/chunks.json."""
    if not settings.INDEX_FILE.exists():
        return []
    try:
        with open(settings.INDEX_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.error(f"Error loading chunks index from {settings.INDEX_FILE}: {e}")
        return []

def save_chunks_index(chunks: List[Dict[str, Any]]) -> None:
    """Saves chunks list to data/index/chunks.json."""
    settings.INDEX_FILE.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(settings.INDEX_FILE, "w", encoding="utf-8") as f:
            json.dump(chunks, f, indent=2, ensure_ascii=False)
        logger.info(f"Saved {len(chunks)} chunks to {settings.INDEX_FILE}")
    except Exception as e:
        logger.error(f"Error saving chunks index: {e}")

def add_chunks_to_index(new_chunks: List[Dict[str, Any]]) -> None:
    """Appends new chunks to the index file, omitting raw vector floats for storage efficiency."""
    existing = load_chunks_index()
    # Strip vector embeddings before saving to JSON index to keep it lightweight for BM25
    clean_chunks = []
    for c in new_chunks:
        clean_chunks.append({
            "chunk_id": c.get("chunk_id"),
            "document_id": c.get("document_id"),
            "document_name": c.get("document_name"),
            "page_or_section": c.get("page_or_section"),
            "text": c.get("text")
        })
    existing.extend(clean_chunks)
    save_chunks_index(existing)

def remove_document_from_index(document_id: str) -> None:
    """Removes all chunks of a given document_id from the JSON index."""
    existing = load_chunks_index()
    filtered = [c for c in existing if c.get("document_id") != document_id]
    save_chunks_index(filtered)

# ----------------- Text Extractors -----------------

def extract_text_from_pdf(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts text page-by-page from PDF with OCR fallback for scanned pages."""
    doc = fitz.open(str(file_path))
    pages_data = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()

        # If page text is very sparse, attempt OCR via pytesseract if available
        if len(text) < 40 and pytesseract is not None:
            try:
                pix = page.get_pixmap(dpi=150)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_text = pytesseract.image_to_string(img).strip()
                if ocr_text:
                    text = f"[OCR extracted]\n{ocr_text}"
            except Exception as ocr_err:
                logger.warning(f"OCR fallback failed on PDF page {page_num + 1}: {ocr_err}")

        if text:
            pages_data.append({
                "page_or_section": f"Page {page_num + 1}",
                "text": text
            })

    doc.close()
    return pages_data

def extract_text_from_docx(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts text and table rows from Word DOCX."""
    doc = docx.Document(str(file_path))
    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                full_text.append(row_text)

    combined = "\n\n".join(full_text)
    if not combined:
        return []
    return [{"page_or_section": "Document Body", "text": combined}]

def extract_text_from_txt(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts text from plain text file."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read().strip()
    return [{"page_or_section": "Full Text", "text": text}] if text else []

def extract_text_from_csv(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts and formats rows from CSV file."""
    lines = []
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            clean_row = [val.strip() for val in row if val.strip()]
            if clean_row:
                lines.append(f"Row {i+1}: " + " | ".join(clean_row))

    text = "\n".join(lines)
    return [{"page_or_section": "CSV Data", "text": text}] if text else []

def extract_text_from_xlsx(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts text from Excel spreadsheets per worksheet."""
    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    sheets_data = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sheet_lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(val).strip() for val in row if val is not None and str(val).strip()]
            if cells:
                sheet_lines.append(" | ".join(cells))

        if sheet_lines:
            sheets_data.append({
                "page_or_section": f"Sheet: {sheet_name}",
                "text": "\n".join(sheet_lines)
            })

    wb.close()
    return sheets_data

def extract_text_from_pptx(file_path: Path) -> List[Dict[str, Any]]:
    """Extracts text slide-by-slide from PowerPoint presentations."""
    prs = pptx.Presentation(str(file_path))
    slides_data = []

    for i, slide in enumerate(prs.slides):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_lines.append(line)

        if slide_lines:
            slides_data.append({
                "page_or_section": f"Slide {i + 1}",
                "text": "\n".join(slide_lines)
            })

    return slides_data

def extract_text_from_image(file_path: Path) -> List[Dict[str, Any]]:
    """Performs OCR text extraction on image formats."""
    if pytesseract is None:
        logger.warning("pytesseract is not available for image OCR.")
        return []
    try:
        img = Image.open(str(file_path))
        text = pytesseract.image_to_string(img).strip()
        return [{"page_or_section": "Image OCR", "text": text}] if text else []
    except Exception as e:
        logger.error(f"Image OCR extraction error on {file_path.name}: {e}")
        return []

def extract_document_content(file_path: Path) -> List[Dict[str, Any]]:
    """Dispatches file extraction based on extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    elif suffix == ".docx":
        return extract_text_from_docx(file_path)
    elif suffix == ".txt":
        return extract_text_from_txt(file_path)
    elif suffix == ".csv":
        return extract_text_from_csv(file_path)
    elif suffix == ".xlsx":
        return extract_text_from_xlsx(file_path)
    elif suffix == ".pptx":
        return extract_text_from_pptx(file_path)
    elif suffix in (".png", ".jpg", ".jpeg", ".webp"):
        return extract_text_from_image(file_path)
    else:
        logger.warning(f"Unsupported file format: {suffix}")
        return []

# ----------------- Ingestion Pipeline -----------------

def ingest_document(
    file_path: Path,
    original_filename: Optional[str] = None,
    document_id: Optional[str] = None
) -> Dict[str, Any]:
    """
    Complete ingestion pipeline:
    1. Extracts text from file with metadata.
    2. Chunks content using RecursiveCharacterTextSplitter.
    3. Generates 384-d dense embeddings.
    4. Indexes vectors into Milvus.
    5. Persists chunks to data/index/chunks.json for BM25 retrieval.
    """
    doc_id = document_id or str(uuid.uuid4())[:8]
    doc_name = original_filename or file_path.name

    logger.info(f"Starting ingestion for document '{doc_name}' (ID: {doc_id})...")

    # 1. Text Extraction
    raw_sections = extract_document_content(file_path)
    if not raw_sections:
        logger.warning(f"No extractable text found in '{doc_name}'.")
        return {
            "document_id": doc_id,
            "document_name": doc_name,
            "chunk_count": 0,
            "status": "warning",
            "message": "No text could be extracted from the document."
        }

    # 2. Text Chunking
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""]
    )

    chunks: List[Dict[str, Any]] = []
    chunk_counter = 0

    for section in raw_sections:
        section_text = section["text"]
        section_label = section["page_or_section"]
        sub_chunks = splitter.split_text(section_text)

        for sub_chunk in sub_chunks:
            if not sub_chunk.strip():
                continue
            chunk_counter += 1
            chunks.append({
                "chunk_id": f"{doc_id}_{chunk_counter}",
                "document_id": doc_id,
                "document_name": doc_name,
                "page_or_section": section_label,
                "text": sub_chunk.strip()
            })

    if not chunks:
        return {
            "document_id": doc_id,
            "document_name": doc_name,
            "chunk_count": 0,
            "status": "warning",
            "message": "Content resulted in 0 usable chunks."
        }

    # 3. Dense Embedding Generation
    texts = [c["text"] for c in chunks]
    vectors = embed_documents(texts)
    for c, vec in zip(chunks, vectors):
        c["vector"] = vec

    # 4. Milvus Vector Storage
    milvus_store.insert_chunks(chunks)

    # 5. Persistent Local Index for BM25
    add_chunks_to_index(chunks)

    logger.info(f"Successfully indexed document '{doc_name}' with {len(chunks)} chunks.")
    return {
        "document_id": doc_id,
        "document_name": doc_name,
        "chunk_count": len(chunks),
        "status": "success",
        "message": f"Successfully ingested {len(chunks)} chunks."
    }
